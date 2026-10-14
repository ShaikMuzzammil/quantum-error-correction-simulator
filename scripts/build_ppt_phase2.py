#!/usr/bin/env python3
"""Append Phase 2 slides (8-12) to Quantum_error_correction_plain.pptx.
Matches the existing design: white bg, Georgia bold #1B2A49 titles, Calibri body,
gray footer + page number. Saves as Quantum_error_correction_Phase2.pptx."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = '/home/z/my-project/upload/Quantum_error_correction_plain.pptx'
OUT = '/home/z/my-project/download/Quantum_error_correction_Phase2.pptx'
ASSETS = '/home/z/my-project/download/ppt-assets'

NAVY = RGBColor(0x1B, 0x2A, 0x49)
BODY = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x6B, 0x72, 0x80)
MID = RGBColor(0x37, 0x41, 0x51)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
BORDER = RGBColor(0xD1, 0xD5, 0xDB)
VIOLET = RGBColor(0x6D, 0x28, 0xD9)
VIOLET_BG = RGBColor(0xEE, 0xF2, 0xFF)
VIOLET_BD = RGBColor(0xC7, 0xD2, 0xFE)
CYAN = RGBColor(0x0E, 0x74, 0x90)
CYAN_BG = RGBColor(0xEC, 0xFE, 0xFF)
CYAN_BD = RGBColor(0xA5, 0xF3, 0xFC)
GREEN = RGBColor(0x04, 0x78, 0x57)
GREEN_BG = RGBColor(0xEC, 0xFD, 0xF5)
GREEN_BD = RGBColor(0xA7, 0xF3, 0xD0)
AMBER_BG = RGBColor(0xFF, 0xFB, 0xEB)
AMBER_BD = RGBColor(0xFD, 0xE6, 0x8A)
AMBER = RGBColor(0x92, 0x40, 0x0E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB9, 0x1C, 0x1C)

prs = Presentation(SRC)
# find DEFAULT layout
layout = None
for ly in prs.slide_layouts:
    if ly.name == 'DEFAULT':
        layout = ly
        break
if layout is None:
    layout = prs.slide_layouts[0]

def add_slide():
    return prs.slides.add_slide(layout)

def no_shadow(sh):
    try:
        sh.shadow.inherit = False
    except Exception:
        pass

def box(slide, l, t, w, h, fill=None, line=None, line_w=0.75, round_=False, radius=0.06):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    if round_:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    no_shadow(shp)
    shp.text_frame.word_wrap = True
    return shp

def txt(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=2, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph = list of (text, font, size, bold, color)."""
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for (text, fname, size, bold, color) in para:
            r = p.add_run()
            r.text = text
            r.font.name = fname
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
    return tb

def title_bar(slide, title, page_no):
    txt(slide, 0.70, 0.40, 11.90, 0.80, [[(title, 'Georgia', 30, True, NAVY)]])
    txt(slide, 0.70, 7.00, 8.00, 0.30,
        [[('Interactive Quantum Error Correction Sim', 'Calibri', 10, False, GRAY)]])
    txt(slide, 12.00, 7.00, 0.60, 0.30, [[(str(page_no), 'Calibri', 10, False, GRAY)]],
        align=PP_ALIGN.RIGHT)

def chip(slide, l, t, w, h, text, bg, bd, fg, size=10.5, bold=False):
    c = box(slide, l, t, w, h, fill=bg, line=bd, line_w=1.0, round_=True, radius=0.5)
    tf = c.text_frame
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = 'Calibri'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = fg
    return c

G = 'Georgia'
C = 'Calibri'

# ============================================================ SLIDE 8
s = add_slide()
title_bar(s, 'Phase 2: What We Built', 8)
txt(s, 0.70, 1.18, 11.90, 0.42,
    [[('Every Phase-1 objective is implemented and working; Phase 2 upgraded the whole prototype into a '
       'premium, self-explaining simulator suite — larger scale, simpler language, real one-shot simulations.',
       C, 12.5, False, MID)]], line_spacing=1.05)

cards = [
    ('1. Aurora-Glass Interface (all 7 tabs)',
     'Layered aurora backdrop, frosted-glass panels, gradient hero, shimmer buttons — a consistent visual '
     'language across Overview, Simulator Lab, 3D Lab, Performance, Learn, Reference and About.', VIOLET_BG, VIOLET_BD, VIOLET),
    ('2. Live Pipeline Playback',
     'Every Simulator Lab shot plays its six stages — Encode → Noise → Syndrome → Correct → Decode → '
     'Measure — on an animated rail with an active glow, then rests at "complete". Replay with one click.', CYAN_BG, CYAN_BD, CYAN),
    ('3. Streaming Monte Carlo + Code Showdown',
     'Performance-Lab estimates converge LIVE while thousands of shots run; one click races all 9 QEC codes '
     'under identical noise and ranks them (gold/silver/bronze) with a full statistics table.', GREEN_BG, GREEN_BD, GREEN),
    ('4. 3D/4D Lab Rebuilt — 4 Live Simulators',
     'Bloch Flight 4D (target pin + fidelity chart), State Cityscape 3D (8-32 shaded towers on a real street '
     'grid), Spacetime Lattice 4D (d up to 9 × 16 rounds, connected time axis) — all one-shot runs.', VIOLET_BG, VIOLET_BD, VIOLET),
    ('5. Threshold Terrain 4D (new, large-scale)',
     'Thousands of complete surface-code experiments run live in the browser (2,520-4,200 per run) and grow a '
     '3-D terrain of measured logical error rates — the famous threshold, measured in front of you.', AMBER_BG, AMBER_BD, AMBER),
    ('6. Understanding Layer + Deploy Guide',
     '? explainers on every input, Results / How-computed / What-it-means tabs, PDF / JSON / CSV exports, '
     'plain-English rewrite of all lab text, and a step-by-step GitHub → Vercel deploy guide in the README.', CYAN_BG, CYAN_BD, CYAN),
]
pos = [(0.70, 1.78), (4.83, 1.78), (8.96, 1.78), (0.70, 3.85), (4.83, 3.85), (8.96, 3.85)]
for (t, body, bg, bd, fg), (l, tp) in zip(cards, pos):
    card = box(s, l, tp, 3.67, 1.92, fill=bg, line=bd, line_w=1.0, round_=True, radius=0.055)
    tf = card.text_frame
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.09)
    tf.margin_bottom = Inches(0.06)
    p1 = tf.paragraphs[0]
    r = p1.add_run(); r.text = t
    r.font.name = C; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = NAVY
    p2 = tf.add_paragraph(); p2.space_before = Pt(3); p2.line_spacing = 1.02
    r2 = p2.add_run(); r2.text = body
    r2.font.name = C; r2.font.size = Pt(9.5); r2.font.color.rgb = MID

badge = box(s, 0.70, 6.05, 11.93, 0.62, fill=GREEN_BG, line=GREEN_BD, line_w=1.0, round_=True, radius=0.5)
tf = badge.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = 'Engineering guarantees:  125 / 125 engine verification tests green  ·  100% client-side (no backend, no env vars)  ·  one-shot simulation discipline — every run plays once and stops'
r.font.name = C; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = GREEN

# ============================================================ SLIDE 9 — architecture diagram
s = add_slide()
title_bar(s, 'System Architecture — Phase 2 Build', 9)

def layer(y, h, label, label_color, chips, chip_bg, chip_bd, chip_fg):
    lab = box(s, 0.70, y, 2.45, h, fill=NAVY, line=None, round_=True, radius=0.10)
    tf = lab.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.10); tf.margin_right = Inches(0.06)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.name = C; r.font.size = Pt(11.5); r.font.bold = True; r.font.color.rgb = WHITE
    # chips area
    n = len(chips)
    total_w = 9.30
    gap = 0.12
    cw = (total_w - gap * (n - 1)) / n
    x = 3.33
    for ch in chips:
        chip(s, x, y + 0.10, cw, h - 0.20, ch, chip_bg, chip_bd, chip_fg, size=9.5, bold=True)
        x += cw + gap

# Layer 1 — interface
layer(1.30, 0.92, 'INTERFACE\n7 guided tabs',
      NAVY,
      ['Overview\nworked example', 'Simulator Lab\n9 codes × 8 noises', '3D Quantum Lab\n4 live simulators',
       'Performance Lab\nlive Monte Carlo', 'Learn + Reference\nzero-to-QEC course'],
      VIOLET_BG, VIOLET_BD, NAVY)
arrow(1) if False else None
def down_arrow(y):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.42), Inches(y), Inches(0.5), Inches(0.28))
    a.fill.solid(); a.fill.fore_color.rgb = RGBColor(0xA5, 0xB4, 0xFC)
    a.line.fill.background(); no_shadow(a)

down_arrow(2.28)
# Layer 2 — visualization components
layer(2.62, 0.92, 'VISUALIZATION\ncomponents',
      NAVY,
      ['Canvas-3D engine\norbit cam · depth-sorted bars', 'Recharts\nlive convergence charts',
       'SVG renderers\ncircuit · Bloch · lattice', 'HUD overlays\nround · fidelity · verdicts',
       'Report builders\nPDF · JSON · CSV'],
      CYAN_BG, CYAN_BD, NAVY)
down_arrow(3.60)
# Layer 3 — engine
layer(3.94, 0.92, 'QEC ENGINE\npure TypeScript\n125-test verified',
      NAVY,
      ['engine.ts\nstate vector (2⁹ amplitudes)', 'noise.ts\n8 quantum channels',
       'codes.ts\n9 QEC codes + tables', 'pipeline.ts\nshots · MC · sweeps',
       'surface.ts\nsurface code + MWPM'],
      GREEN_BG, GREEN_BD, NAVY)
down_arrow(4.92)
# Layer 4 — outputs
layer(5.26, 0.86, 'OUTPUTS\nexportable',
      NAVY,
      ['Printable PDF report\ninputs · outputs · methodology', 'JSON\nfull machine-readable data',
       'CSV\nper-round / per-point data', 'Live HUDs\nstats while running'],
      LIGHT, BORDER, BODY)

note = box(s, 0.70, 6.30, 11.93, 0.52, fill=AMBER_BG, line=AMBER_BD, line_w=1.0, round_=True, radius=0.5)
tf = note.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = 'Stack: Next.js 16 (App Router) · React 19 · TypeScript 5 · Tailwind CSS 4 — 100% client-side: no backend, no database, no environment variables. Deploy = GitHub push → Vercel.'
r.font.name = C; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = AMBER

# ============================================================ SLIDE 10 — comparative analysis
s = add_slide()
title_bar(s, 'Comparative Analysis — Phase 1 vs Phase 2', 10)

rows = [
    ('Area', 'Phase-1 objective (proposal)', 'Phase-2 delivered (this build)', 'Status'),
    ('Noise simulation', 'Simulate multiple quantum noise models',
     '8 channels: bit-flip, phase-flip, depolarizing, amplitude/phase damping, thermal relaxation (T1/T2), coherent rotation, crosstalk — exact quantum-trajectory sampling', 'Exceeded'),
    ('Performance evaluation', 'Compare error rates with and without QEC',
     'Streaming Monte Carlo up to 100,000 shots, noise sweeps with live charts, and Code Showdown ranking all 9 codes by improvement factor', 'Exceeded'),
    ('Configurable QEC', 'A configurable environment for QEC experimentation',
     '9 code configurations ([[3]] to Shor [[9]] and surface d=3-9), manual per-qubit X/Z/Y error injection, RNG seed control, custom input states', 'Exceeded'),
    ('Interface & usability', '— (added during Phase 2)',
     '7 guided tabs, ?-explainers on every input, Results / How-computed / What-it-means tabs, plain-language rewrite, dark aurora-glass theme', 'New'),
    ('Scale of simulation', '— (added during Phase 2)',
     '4 live 3D/4D simulators; Threshold Terrain runs 2,520-4,200 complete surface-code experiments per run, live in the browser', 'New'),
    ('Correctness', '— (added during Phase 2)',
     '125-test engine suite: 3-qubit fidelity matches 1−P(≥2 errors), surface-code P_L(d=3) = 18p² exact, threshold ordering reproduced', 'New'),
]
tbl_shape = s.shapes.add_table(len(rows), 4, Inches(0.70), Inches(1.30), Inches(11.93), Inches(5.35))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(1.75)
tbl.columns[1].width = Inches(3.30)
tbl.columns[2].width = Inches(5.68)
tbl.columns[3].width = Inches(1.20)
tbl.first_row = True
tbl.horz_banding = False
for ri, row in enumerate(rows):
    tbl.rows[ri].height = Inches(0.62 if ri else 0.40)
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.name = C
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
            r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = WHITE
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else RGBColor(0xF6, 0xF7, 0xF9)
            r.font.size = Pt(9.5)
            r.font.bold = (ci == 0) or (ci == 3)
            r.font.color.rgb = NAVY if ci == 0 else (GREEN if ci == 3 else MID)

verdict = box(s, 0.70, 6.78, 11.93, 0.50, fill=GREEN_BG, line=GREEN_BD, line_w=1.0, round_=True, radius=0.5)
tf = verdict.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = 'Verdict: all three Phase-1 objectives are fully implemented — and Phase 2 added the interface, scale and verification layers on top.'
r.font.name = C; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = GREEN

# ============================================================ SLIDE 11 — noise-model simulation photos
s = add_slide()
title_bar(s, 'Noise-Model Simulation in Action', 11)
txt(s, 0.70, 1.16, 11.90, 0.40,
    [[('All three images are real runs of the upgraded prototype: noise hits the code, the decoder answers, '
       'and every run stops by itself when complete.', C, 12, False, MID)]])

imgs = [
    (f'{ASSETS}/final-bloch-flight3.png', 'Bloch Flight 4D — bit-flip noise p=10%, 30 rounds.',
     'QEC holds the logical state exactly on the target pin (F = 100.0%) while the unprotected qubit collapses to F = 49.1%.'),
    (f'{ASSETS}/final-spacetime.png', 'Spacetime Lattice 4D — surface code d=5, 8 syndrome rounds.',
     'Defect stars light up where stabilizers flag errors; the MWPM decoder draws matching chains through spacetime (1 logical fail detected and reset).'),
    (f'{ASSETS}/final-terrain2.png', 'Threshold Terrain 4D — 4,200 live Monte Carlo experiments.',
     'Measured logical error rate for every (noise p, distance d) pair; the verdict line reports the threshold where bigger codes stop helping.'),
]
x = 0.70
for path, cap1, cap2 in imgs:
    pic = s.shapes.add_picture(path, Inches(x), Inches(1.62), width=Inches(3.90))
    # aspect 1440x900 → height = 3.90 * 900/1440 = 2.4375
    txt(s, x, 4.16, 3.90, 0.55, [[(cap1, C, 10, True, NAVY)]])
    txt(s, x, 4.62, 3.90, 1.05, [[(cap2, C, 9.5, False, MID)]], line_spacing=1.02)
    # thin border frame
    fr = box(s, x - 0.015, 1.605, 3.93, 2.47, fill=None, line=BORDER, line_w=1.0)
    x += 4.065

strip = box(s, 0.70, 5.90, 11.93, 0.80, fill=CYAN_BG, line=CYAN_BD, line_w=1.0, round_=True, radius=0.12)
tf = strip.text_frame
tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.line_spacing = 1.05
r = p.add_run()
r.text = ('Noise models available in every lab: ')
r.font.name = C; r.font.size = Pt(10.5); r.font.bold = False; r.font.color.rgb = MID
r2 = p.add_run()
r2.text = 'Bit-Flip (X) · Phase-Flip (Z) · Depolarizing · Amplitude Damping · Phase Damping · Thermal Relaxation (T1/T2) · Coherent Over-Rotation · Crosstalk — the same channels drive the 2D labs and the 3D/4D visualizers.'
r2.font.name = C; r2.font.size = Pt(10.5); r2.font.bold = True; r2.font.color.rgb = CYAN

# ============================================================ SLIDE 12 — results & outputs
s = add_slide()
title_bar(s, 'Results & Outputs — Measured, Not Claimed', 12)

imgs2 = [
    (f'{ASSETS}/final-simulator-lab.png', 'Simulator Lab — a full pipeline shot.',
     'Circuit, syndrome table, Bloch comparison and state inspector: "Quantum information recovered", fidelity 100.00% with QEC.'),
    (f'{ASSETS}/final-showdown.png', 'Code Showdown — all 9 codes raced under identical noise (2,500 shots).',
     '3-Qubit Bit-Flip wins at 4.33× improvement; surface codes and Shor/Steane ranked with full statistics.'),
    (f'{ASSETS}/final-cityscape32c.png', 'State Cityscape 3D — 32 live towers (5-qubit code).',
     'Noise grows error towers, every correction cycle flattens the skyline back to the two codespace towers.'),
]
x = 0.70
for path, cap1, cap2 in imgs2:
    pic = s.shapes.add_picture(path, Inches(x), Inches(1.30), width=Inches(3.90))
    txt(s, x, 3.84, 3.90, 0.55, [[(cap1, C, 10, True, NAVY)]])
    txt(s, x, 4.30, 3.90, 1.0, [[(cap2, C, 9.5, False, MID)]], line_spacing=1.02)
    fr = box(s, x - 0.015, 1.285, 3.93, 2.47, fill=None, line=BORDER, line_w=1.0)
    x += 4.065

stats = [
    ('100.00%', 'logical fidelity recovered\n(Simulator Lab, with QEC)', GREEN),
    ('4.33×', 'best improvement factor\n(Code Showdown, p = 10%)', CYAN),
    ('≈ 8-13%', 'measured threshold — theory: ~10%\n(Threshold Terrain, live MC)', VIOLET),
    ('125 / 125', 'engine verification tests\ngreen before every release', AMBER),
]
x = 0.70
for big, small, color in stats:
    st = box(s, x, 5.42, 2.86, 1.02, fill=WHITE, line=BORDER, line_w=1.0, round_=True, radius=0.10)
    tf = st.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = big
    r.font.name = G; r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = color
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.line_spacing = 0.95
    r2 = p2.add_run(); r2.text = small
    r2.font.name = C; r2.font.size = Pt(8.5); r2.font.color.rgb = MID
    x += 3.025

txt(s, 0.70, 6.60, 11.93, 0.35,
    [[('Deliverable: the prototype ships with a step-by-step GitHub push → Vercel deploy guide (README) — live in ~5 minutes, no configuration needed.',
       C, 10.5, True, NAVY)]], align=PP_ALIGN.CENTER)

prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides.__iter__.__self__._sldIdLst))
