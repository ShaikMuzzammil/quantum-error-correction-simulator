#!/usr/bin/env python3
"""Extract all text content from the Quantum Error Correction PPTX."""
from pptx import Presentation

path = "/home/z/my-project/upload/Quantum_error_correction_plain.pptx"
prs = Presentation(path)

print(f"Total slides: {len(prs.slides)}")
print("=" * 80)

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n{'='*80}\nSLIDE {idx}  (layout: {slide.slide_layout.name})")
    print("-" * 80)
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                print(f"[TEXT] {txt}")
        if shape.has_table:
            tbl = shape.table
            print("[TABLE]")
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                print("  | " + " | ".join(cells))
        if shape.has_chart:
            print(f"[CHART] type={shape.chart.chart_type}")
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            print(f"[NOTES] {notes}")
